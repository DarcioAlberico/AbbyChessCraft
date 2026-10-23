"""O DOCX do AbbyChessCraft (SPEC §12.1, passo J): o escritor próprio, sem o do core.

Lê as mesmas `core.pagina.PaginaExtraida` que o EPUB e escreve pelo python-docx:

- **Títulos** nos estilos "Heading 1" e "Heading 2" (o painel de navegação do Word e o
  sumário que ele monta saem deles); **negrito e itálico** pelos trechos do parágrafo; os
  **símbolos de xadrez** na fonte dos símbolos (`core.exportar.fonte_dos_simbolos`).
- **Diagramas** como imagem, com o texto alternativo do EPUB (o FEN e o lado com a
  procedência, ou que o diagrama não foi lido) no `descr` da figura; imagens iguais são uma
  parte só (o python-docx já guarda pela soma).
- **Páginas impressas**: cada página começa com um marcador (`pg23`), e as referências do
  texto ("página 23", `epub.referencias`) viram hiperlink interno para ele, quando a página
  está no arquivo. Os diagramas numerados pela legenda e os capítulos numerados pelo título
  (`epub.numerados`) têm marcador (`dg6`, `cap3`), e "ver o Diagrama 6" e "no capítulo 3"
  viram hiperlink para eles.
- **Metadados** (`epub.Metadados`): título, autores, assuntos, idioma e identificador nas
  propriedades; o idioma também no estilo base (`w:lang`), senão o Word revisa um livro
  inglês em português.
- **Os mesmos bytes** com o mesmo `SOURCE_DATE_EPOCH` (CA-12.6): as datas das
  propriedades e das entradas do zip vêm de `core.exportar.instante`, e o zip é regravado
  na ordem e com os atributos fixos.

`conferir` relê o arquivo: o pacote abre, o XML do documento é legível, cada imagem citada
existe, cada hiperlink interno tem o seu marcador e os marcadores não se repetem.
"""

from __future__ import annotations

import datetime
import io
import os
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple
from xml.etree import ElementTree

from core.pagina import Figura, PaginaExtraida, Paragrafo

from abbychess import epub

#: A largura do diagrama sem medida de casa e a máxima de qualquer figura, em centímetros.
LARGURA_DO_DIAGRAMA_CM = 6.5
LARGURA_MAXIMA_DO_RECORTE_CM = 9.0
#: O tamanho da casa do diagrama, em pontos: a largura vem de `Figura.casas_de_largura`.
CORPO_DA_CASA_PT = 16.0
#: O desenho da página, o do escritor do core (F111): Georgia 11, justificado, recuo de
#: primeira linha, margens de 2,2 cm.
FONTE_DO_TEXTO, CORPO_DO_TEXTO_PT, MARGEM_DA_PAGINA_CM = "Georgia", 11.0, 2.2
#: A cor do hiperlink interno (a do estilo "Hyperlink" do Word).
COR_DO_LINK = "0563C1"

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


@dataclass
class EscritoDocx:
    """O que o DOCX levou: para o relatório da exportação (§12.2)."""

    paginas: int = 0
    paragrafos: int = 0
    titulos: int = 0
    figuras: int = 0
    links: int = 0
    #: Os parágrafos com algum trecho realçado (`realce`).
    realcados: int = 0


# ----------------------------------------------------------------------
# Pedaços de OOXML
# ----------------------------------------------------------------------

def _desenho_da_pagina(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    normal = doc.styles["Normal"]
    normal.font.name = FONTE_DO_TEXTO
    normal.font.size = Pt(CORPO_DO_TEXTO_PT)
    formato = normal.paragraph_format
    formato.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    formato.first_line_indent = Pt(CORPO_DO_TEXTO_PT * 1.2)
    formato.space_after = Pt(CORPO_DO_TEXTO_PT * 0.35)
    for secao in doc.sections:
        secao.top_margin = secao.bottom_margin = Cm(MARGEM_DA_PAGINA_CM)
        secao.left_margin = secao.right_margin = Cm(MARGEM_DA_PAGINA_CM)


def _idioma_no_estilo(doc, idioma: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    for velho in rpr.findall(qn("w:lang")):
        rpr.remove(velho)
    lang = OxmlElement("w:lang")
    for atributo in ("w:val", "w:eastAsia", "w:bidi"):
        lang.set(qn(atributo), idioma)
    rpr.append(lang)


def _marcador(paragrafo, nome: str, ident: int) -> None:
    """Um marcador do Word (`w:bookmarkStart`/`End`) envolvendo o parágrafo."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    inicio = OxmlElement("w:bookmarkStart")
    inicio.set(qn("w:id"), str(ident))
    inicio.set(qn("w:name"), nome)
    fim = OxmlElement("w:bookmarkEnd")
    fim.set(qn("w:id"), str(ident))
    p = paragrafo._p
    ppr = p.pPr
    if ppr is not None:
        ppr.addnext(inicio)
    else:
        p.insert(0, inicio)
    p.append(fim)


def _texto_com_estilo(paragrafo, bloco: Paragrafo, simbolos: str, familia: str,
                      links: Sequence[Tuple[int, int, str]], quebras: bool = False,
                      realce: bool = False) -> None:
    """Os trechos do parágrafo como runs: negrito, itálico, fonte dos símbolos e o
    hiperlink interno por fora (o mesmo corte de `epub.xhtml_do_texto`)."""
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    # Como no EPUB: o comprimento fica; a quebra de linha é espaço, como no XHTML, salvo com
    # `quebras` (as linhas impressas mantidas): aí ela fica, e o python-docx a escreve `w:br`.
    texto = epub.sem_proibidos(bloco.texto).replace("\r", " ")
    if not quebras:
        texto = texto.replace("\n", " ")
    marcas = [[False, False, c in simbolos, -1, False] for c in texto]
    duvidosos = getattr(bloco, "duvidosos", ()) if realce else ()
    for k, trechos in ((0, bloco.negrito), (1, bloco.italico), (4, duvidosos)):
        for inicio, fim in trechos or ():
            for i in range(max(0, inicio), min(fim, len(texto))):
                marcas[i][k] = True
    for n, (inicio, fim, _alvo) in enumerate(links):
        for i in range(max(0, inicio), min(fim, len(texto))):
            marcas[i][3] = n
    hiperlink, do_hiperlink = None, -1
    i = 0
    while i < len(texto):
        j = i
        while j < len(texto) and marcas[j] == marcas[i]:
            j += 1
        negrito, italico, simbolo, link, duvidoso = marcas[i]
        run = paragrafo.add_run(texto[i:j])
        run.bold = negrito or None
        run.italic = italico or None
        if duvidoso:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if simbolo and familia:
            run.font.name = familia
            fontes = run._r.get_or_add_rPr().get_or_add_rFonts()
            fontes.set(qn("w:cs"), familia)
            fontes.set(qn("w:eastAsia"), familia)
        if link >= 0:
            if hiperlink is None or do_hiperlink != link:
                hiperlink = OxmlElement("w:hyperlink")
                hiperlink.set(qn("w:anchor"), links[link][2])
                hiperlink.set(qn("w:history"), "1")
                paragrafo._p.append(hiperlink)
                do_hiperlink = link
            run.font.underline = True
            run.font.color.rgb = RGBColor.from_string(COR_DO_LINK)
            hiperlink.append(run._r)
        else:
            hiperlink, do_hiperlink = None, -1
        i = j


def _figura(doc, figura: Figura, alt: str, numero: int) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    from core.exportar import largura_em_pt

    em_pt = largura_em_pt(figura, CORPO_DA_CASA_PT)
    if em_pt:
        largura = min(Pt(em_pt), Cm(LARGURA_MAXIMA_DO_RECORTE_CM))
    elif figura.fen:
        largura = Cm(LARGURA_DO_DIAGRAMA_CM)
    else:
        proporcao = (figura.largura / figura.altura) if figura.altura else 1.0
        largura = Cm(min(LARGURA_MAXIMA_DO_RECORTE_CM,
                         LARGURA_DO_DIAGRAMA_CM * max(proporcao, 0.5)))
    doc.add_picture(io.BytesIO(figura.png), width=largura)
    paragrafo = doc.paragraphs[-1]
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for propriedades in paragrafo._p.xpath(".//wp:docPr"):
        propriedades.set("descr", epub.limpo(alt))
        propriedades.set("name", f"Diagrama {numero}")


def _ressalva(doc, figura: Figura, idioma: str) -> bool:
    """A legenda visível do diagrama cuja posição ninguém conferiu (PD-21): de quem é a vez
    e "não revisado". Fica no `descr` também, mas quem vê a página precisa ver isto."""
    if not getattr(figura, "revisao_pendente", False):
        return False
    msg = epub.mensagens(idioma)
    partes = [epub.marca_do_lado(figura, idioma)] if figura.fen else []
    partes.append(msg["nao_revisado"])
    paragrafo = doc.add_paragraph(epub.limpo("; ".join(partes)))
    try:
        paragrafo.style = doc.styles["Caption"]
    except KeyError:                            # o modelo sem o estilo: fica no Normal
        pass
    return True


# ----------------------------------------------------------------------
# Escrever
# ----------------------------------------------------------------------

def escrever(paginas: Sequence[PaginaExtraida], caminho: str | Path, metadados: epub.Metadados,
             *, numeros: Optional[Dict[int, int]] = None, quebras: bool = False,
             realce: bool = False) -> EscritoDocx:
    """Escreve o DOCX em `caminho` (ao lado e troca no fim: a falha não estraga o anterior).

    `numeros` é o número impresso de cada página (índice -> número): as referências do
    texto a ele viram hiperlink para o marcador da página."""
    from docx import Document

    from core.exportar import data_no_zip, fonte_dos_simbolos, instante

    if not any(isinstance(b, Figura) or (isinstance(b, Paragrafo) and b.texto.strip())
               for p in paginas for b in p.blocos):
        raise ValueError("Nada para escrever no DOCX: as páginas não têm texto nem diagrama.")
    idioma = metadados.idioma or "pt-BR"
    msg = epub.mensagens(idioma)
    escrito = EscritoDocx(paginas=len(paginas))
    idioma = epub.limpo(idioma)
    doc = Document()
    _desenho_da_pagina(doc)
    _idioma_no_estilo(doc, idioma)

    recurso = fonte_dos_simbolos("".join(b.texto for p in paginas for b in p.blocos
                                         if isinstance(b, Paragrafo)))
    familia, simbolos = (recurso[0], recurso[2]) if recurso else ("", "")

    # Só as páginas que têm o que escrever ganham marcador; só elas são alvo de link.
    com_conteudo = [p for p in paginas if any(
        isinstance(b, Figura) or (isinstance(b, Paragrafo) and b.texto.strip())
        for b in p.blocos)]
    alvos = {}
    for pagina in com_conteudo:
        if numeros and pagina.numero in numeros:
            alvos[numeros[pagina.numero]] = f"pg{pagina.numero + 1}"
    # Os diagramas e os capítulos numerados ganham marcador (`dg6`, `cap3`) e são alvo de
    # "ver o Diagrama 6" e "no capítulo 3", como no EPUB.
    contados = epub.numerados(com_conteudo)
    por_diagrama = {n: f"dg{n}" for n in contados.figuras.values()}
    por_capitulo = {n: f"cap{n}" for n in contados.capitulos.values()}

    ident = 0
    for numero_da_pagina, pagina in enumerate(com_conteudo):
        primeiro = True
        for bloco in pagina.blocos:
            if isinstance(bloco, Figura):
                escrito.figuras += 1
                _figura(doc, bloco, epub.alternativo(bloco, msg, idioma), escrito.figuras)
                paragrafo = doc.paragraphs[-1]
                if primeiro and numero_da_pagina:
                    paragrafo.paragraph_format.page_break_before = True
                if primeiro:
                    ident += 1
                    _marcador(paragrafo, f"pg{pagina.numero + 1}", ident)
                    primeiro = False
                if id(bloco) in contados.figuras:
                    ident += 1
                    _marcador(paragrafo, f"dg{contados.figuras[id(bloco)]}", ident)
                _ressalva(doc, bloco, idioma)
                continue
            elif isinstance(bloco, Paragrafo) and bloco.texto.strip():
                if bloco.titulo:
                    estilo = "Heading 1" if bloco.nivel == 1 else "Heading 2"
                    escrito.titulos += 1
                    links: list = []
                else:
                    estilo = None
                    links = epub.referencias_internas(
                        bloco.texto, idioma, alvos,
                        {} if id(bloco) in contados.legendas else por_diagrama, por_capitulo)
                    escrito.links += len(links)
                paragrafo = doc.add_paragraph(style=estilo)
                _texto_com_estilo(paragrafo, bloco, simbolos, familia, links, quebras, realce)
                if realce and getattr(bloco, "duvidosos", None):
                    escrito.realcados += 1
                escrito.paragrafos += 1
                if id(bloco) in contados.capitulos:
                    ident += 1
                    _marcador(paragrafo, f"cap{contados.capitulos[id(bloco)]}", ident)
            else:
                continue
            if primeiro:
                # A página impressa começa numa página nova, como no livro (o DOCX é o
                # formato de editar o livro, e o número dela fica no marcador).
                if numero_da_pagina:
                    paragrafo.paragraph_format.page_break_before = True
                ident += 1
                _marcador(paragrafo, f"pg{pagina.numero + 1}", ident)
                primeiro = False

    momento = instante()
    data = datetime.datetime.fromtimestamp(momento, datetime.timezone.utc).replace(tzinfo=None)
    propriedades = doc.core_properties
    propriedades.title = epub.limpo(metadados.titulo)
    propriedades.author = epub.limpo("; ".join(metadados.autores))
    propriedades.subject = epub.limpo("; ".join(metadados.assuntos))
    propriedades.keywords = epub.limpo("; ".join(metadados.assuntos))
    propriedades.language = idioma
    propriedades.identifier = epub.limpo(metadados.identificador_do_livro())
    propriedades.last_modified_by = "AbbyChessCraft"
    propriedades.comments = ""
    propriedades.revision = 1
    propriedades.created = data
    propriedades.modified = data

    bruto = io.BytesIO()
    doc.save(bruto)
    partes = _partes_finais(bruto.getvalue(), recurso)
    _gravar_reprodutivel(Path(caminho), partes, data_no_zip(momento))
    return escrito


def _partes_finais(dados: bytes, recurso) -> Dict[str, bytes]:
    """As partes do pacote como vão para o disco: sem a miniatura e as sobras do modelo
    do python-docx (a página em branco, a data de impressão), e com a fonte dos símbolos
    embutida do jeito que o Word espera (ECMA-376 §15.2.13: ofuscada, o `.odttf` no
    `[Content_Types]`, a fonte no `fontTable` com o relacionamento, e o
    `w:embedTrueTypeFonts` nas configurações, logo depois do `w:zoom`)."""
    import re

    from core.exportar import GUID_DA_FONTE, REL_DE_FONTE, TIPO_OFUSCADO, ofuscar

    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        partes = {info.filename: z.read(info) for info in z.infolist()}
    miniatura = "docProps/thumbnail.jpeg"
    if miniatura in partes:
        del partes[miniatura]
        partes["_rels/.rels"] = re.sub(rb'<Relationship [^>]*Target="docProps/thumbnail\.jpeg"/>',
                                       b"", partes["_rels/.rels"])
        partes["[Content_Types].xml"] = re.sub(
            rb'<Override PartName="/docProps/thumbnail\.jpeg"[^>]*/>', b"",
            partes["[Content_Types].xml"])
    if "docProps/core.xml" in partes:
        partes["docProps/core.xml"] = re.sub(rb"<cp:lastPrinted>.*?</cp:lastPrinted>", b"",
                                             partes["docProps/core.xml"])
    if not recurso:
        return partes
    familia, origem = recurso[0], recurso[1]
    with open(origem, "rb") as f:
        partes["word/fonts/fonte1.odttf"] = ofuscar(f.read())
    tipos = partes["[Content_Types].xml"].decode("utf-8")
    if "odttf" not in tipos:
        abertura = tipos.index(">", tipos.index("<Types")) + 1
        tipos = (tipos[:abertura] + f'<Default Extension="odttf" ContentType="{TIPO_OFUSCADO}"/>'
                 + tipos[abertura:])
        partes["[Content_Types].xml"] = tipos.encode("utf-8")
    nome = familia.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;")
    tabela = partes["word/fontTable.xml"].decode("utf-8")
    corte = tabela.index("</w:fonts>")
    tabela = (tabela[:corte] + f'<w:font w:name="{nome}"><w:embedRegular r:id="rIdFonte1" '
              f'w:fontKey="{GUID_DA_FONTE}"/></w:font>' + tabela[corte:])
    if "xmlns:r=" not in tabela[:tabela.index(">", tabela.index("<w:fonts"))]:
        tabela = tabela.replace("<w:fonts ", f'<w:fonts xmlns:r="{_R}" ', 1)
    partes["word/fontTable.xml"] = tabela.encode("utf-8")
    partes["word/_rels/fontTable.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rIdFonte1" Type="{REL_DE_FONTE}" Target="fonts/fonte1.odttf"/>'
        "</Relationships>").encode("utf-8")
    ajustes = partes["word/settings.xml"].decode("utf-8")
    if "embedTrueTypeFonts" not in ajustes:
        if "<w:zoom" in ajustes:
            corte = ajustes.index("/>", ajustes.index("<w:zoom")) + 2
        else:
            corte = ajustes.index(">", ajustes.index("<w:settings")) + 1
        ajustes = ajustes[:corte] + "<w:embedTrueTypeFonts/>" + ajustes[corte:]
        partes["word/settings.xml"] = ajustes.encode("utf-8")
    return partes


def _gravar_reprodutivel(caminho: Path, partes: Dict[str, bytes], data_hora) -> None:
    """As partes num zip com a mesma data e os mesmos atributos em toda entrada (o
    python-docx carimba a hora de agora), ao lado; conferido antes de trocar com o
    anterior: um DOCX que não passa em `conferir` não é publicado."""
    caminho.parent.mkdir(parents=True, exist_ok=True)
    fd, temporario = tempfile.mkstemp(prefix=".docx-", suffix=".tmp", dir=caminho.parent)
    os.close(fd)
    try:
        with zipfile.ZipFile(temporario, "w", zipfile.ZIP_DEFLATED) as destino:
            for nome, dados in partes.items():
                info = zipfile.ZipInfo(nome, date_time=data_hora)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0o644 << 16
                destino.writestr(info, dados)
        problemas = conferir(temporario)
        if problemas:
            raise ValueError("O DOCX não passou na conferência: " + "; ".join(problemas))
        os.chmod(temporario, 0o644)
        os.replace(temporario, caminho)
    except BaseException:
        Path(temporario).unlink(missing_ok=True)
        raise


# ----------------------------------------------------------------------
# Conferir
# ----------------------------------------------------------------------

def conferir(caminho: str | Path) -> List[str]:
    """Os problemas do DOCX que o Word recusaria ou mostraria quebrado; lista vazia é
    arquivo são."""
    problemas: List[str] = []
    try:
        z = zipfile.ZipFile(caminho)
    except (zipfile.BadZipFile, OSError) as erro:
        return [f"o DOCX não abre como zip ({erro})"]
    with z:
        nomes = set(z.namelist())
        for obrigatorio in ("[Content_Types].xml", "word/document.xml",
                            "word/_rels/document.xml.rels"):
            if obrigatorio not in nomes:
                problemas.append(f"falta {obrigatorio}")
        if problemas:
            return problemas
        try:
            documento = ElementTree.fromstring(z.read("word/document.xml"))
            relacoes = ElementTree.fromstring(z.read("word/_rels/document.xml.rels"))
        except ElementTree.ParseError as erro:
            return [f"XML ilegível ({erro})"]
        alvos = {r.get("Id"): r.get("Target") for r in relacoes}
        for blip in documento.iter(f"{{{_A}}}blip"):
            ident = blip.get(f"{{{_R}}}embed")
            alvo = alvos.get(ident)
            if alvo is None or f"word/{alvo}" not in nomes:
                problemas.append(f"a imagem {ident} não está no pacote")
        marcadores, ids = set(), set()
        for marca in documento.iter(f"{{{_W}}}bookmarkStart"):
            nome, ident = marca.get(f"{{{_W}}}name"), marca.get(f"{{{_W}}}id")
            if nome in marcadores:
                problemas.append(f"o marcador {nome} se repete")
            if ident in ids:
                problemas.append(f"o id de marcador {ident} se repete")
            marcadores.add(nome)
            ids.add(ident)
        fechados = {m.get(f"{{{_W}}}id") for m in documento.iter(f"{{{_W}}}bookmarkEnd")}
        for ident in ids - fechados:
            problemas.append(f"o marcador de id {ident} não fecha")
        for link in documento.iter(f"{{{_W}}}hyperlink"):
            ancora = link.get(f"{{{_W}}}anchor")
            if ancora is None and link.get(f"{{{_R}}}id") is None:
                problemas.append("hiperlink sem destino")
            elif ancora is not None and ancora not in marcadores:
                problemas.append(f"o hiperlink para {ancora} não tem marcador")
        figuras = [d.get("id") for d in documento.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")]
        if len(figuras) != len(set(figuras)):
            problemas.append("ids de figura repetidos")
        if "word/fontTable.xml" in nomes:
            tabela = z.read("word/fontTable.xml")
            if b"embedRegular" in tabela:
                rels = "word/_rels/fontTable.xml.rels"
                if rels not in nomes:
                    problemas.append("a fonte embutida não tem relacionamento")
                else:
                    for rel in ElementTree.fromstring(z.read(rels)):
                        if f"word/{rel.get('Target')}" not in nomes:
                            problemas.append(f"a fonte {rel.get('Target')} não está no pacote")
    try:
        from docx import Document

        Document(str(caminho))
    except Exception as erro:  # noqa: BLE001 — o que o python-docx não abre, o Word também não
        problemas.append(f"o python-docx não abre o arquivo ({erro})")
    return problemas
