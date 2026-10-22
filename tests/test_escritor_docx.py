"""O escritor próprio de DOCX (passo J): títulos, estilos, diagramas com texto alternativo,
marcadores e hiperlinks de página, metadados, os mesmos bytes e a conferência."""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from PIL import Image

from core.pagina import Figura, PaginaExtraida, Paragrafo

from abbychess import epub
from abbychess import escritor_docx as ed

FEN = "r3k2r/pp3ppp/2n5/3q4/3P4/2N2N2/PP3PPP/R3K2R w - - 0 1"


def _png(cor: int = 0) -> bytes:
    dados = io.BytesIO()
    Image.new("L", (16, 16), cor).save(dados, format="PNG")
    return dados.getvalue()


def _livro():
    return [
        PaginaExtraida(numero=0, blocos=[
            Paragrafo("Capítulo Um", titulo=True, nivel=1),
            Paragrafo("Negrito e itálico 1.♘f3; a solução está na página 12.",
                      negrito=[(0, 7)], italico=[(10, 17)])]),
        PaginaExtraida(numero=1, blocos=[
            Paragrafo("Uma seção", titulo=True, nivel=2),
            Figura(_png(0), 16, 16, fen=FEN, origem="render", lado_a_jogar="b",
                   lado_origem="usuario"),
            Figura(_png(0), 16, 16, fen=FEN, origem="render"),
            Figura(_png(200), 16, 24, origem="recorte", aviso="posição não conferida")]),
        PaginaExtraida(numero=2, blocos=[Paragrafo("A solução.")]),
    ]


def _meta(**campos):
    return epub.Metadados("O Livro", autores=("A. Autor", "B. Autora"),
                          assuntos=("Xadrez", "Finais"), **campos)


def test_o_docx_tem_titulos_estilos_figuras_e_links(tmp_path):
    destino = tmp_path / "livro.docx"
    escrito = ed.escrever(_livro(), destino, _meta(), numeros={0: 10, 1: 11, 2: 12})
    assert ed.conferir(destino) == []
    assert (escrito.titulos, escrito.figuras, escrito.links) == (2, 3, 1)

    doc = Document(str(destino))
    estilos = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert estilos[:2] == [("Heading 1", "Capítulo Um"),
                           ("Normal", "Negrito e itálico 1.♘f3; a solução está na página 12.")]
    assert ("Heading 2", "Uma seção") in estilos
    runs = doc.paragraphs[1].runs
    assert runs[0].text == "Negrito" and runs[0].bold
    assert any(r.text == "itálico" and r.italic for r in runs)

    with zipfile.ZipFile(destino) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        imagens = [n for n in z.namelist() if n.startswith("word/media/")]
    assert 'w:anchor="pg3"' in xml and 'w:name="pg3"' in xml   # página 12 = índice 2
    assert f'descr="Diagrama: FEN {FEN} — pretas a jogar (informado na revisão)"' in xml
    assert "Diagrama não lido; imagem da página impressa (posição não conferida)" in xml
    assert len(imagens) == 2                                   # as duas iguais: uma parte

    propriedades = doc.core_properties
    assert propriedades.title == "O Livro"
    assert propriedades.author == "A. Autor; B. Autora"
    assert propriedades.subject == "Xadrez; Finais" and propriedades.language == "pt-BR"
    assert propriedades.identifier.startswith("urn:uuid:")


def test_o_idioma_vai_ao_estilo_base(tmp_path):
    destino = tmp_path / "en.docx"
    ed.escrever(_livro(), destino, _meta(idioma="en"))
    with zipfile.ZipFile(destino) as z:
        estilos = z.read("word/styles.xml").decode("utf-8")
    assert 'w:val="en"' in estilos


def test_ca_12_6_os_mesmos_bytes(tmp_path, monkeypatch):
    monkeypatch.setenv("SOURCE_DATE_EPOCH", "1700000000")
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    ed.escrever(_livro(), a, _meta(), numeros={0: 10, 1: 11, 2: 12})
    ed.escrever(_livro(), b, _meta(), numeros={0: 10, 1: 11, 2: 12})
    assert a.read_bytes() == b.read_bytes()
    with zipfile.ZipFile(a) as z:
        assert {i.date_time for i in z.infolist()} == {(2023, 11, 14, 22, 13, 20)}


def test_sem_numero_impresso_nao_ha_link(tmp_path):
    destino = tmp_path / "a.docx"
    escrito = ed.escrever(_livro(), destino, _meta())
    assert escrito.links == 0 and ed.conferir(destino) == []


def test_livro_vazio_e_recusado(tmp_path):
    with pytest.raises(ValueError, match="Nada para escrever"):
        ed.escrever([PaginaExtraida(numero=0)], tmp_path / "a.docx", _meta())


def test_caracteres_de_controle_do_ocr_nao_quebram_o_docx(tmp_path):
    destino = tmp_path / "a.docx"
    ed.escrever([PaginaExtraida(numero=0, blocos=[Paragrafo("um\x0c texto\x00",
                                                            negrito=[(0, 2)])])],
                destino, _meta())
    assert ed.conferir(destino) == []
    assert Document(str(destino)).paragraphs[0].text == "um  texto "


def test_conferir_acusa_hiperlink_sem_marcador(tmp_path):
    destino = tmp_path / "a.docx"
    ed.escrever(_livro(), destino, _meta(), numeros={0: 10, 1: 11, 2: 12})
    with zipfile.ZipFile(destino) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    partes["word/document.xml"] = partes["word/document.xml"].replace(b'w:name="pg3"',
                                                                      b'w:name="outro"')
    ruim = tmp_path / "ruim.docx"
    with zipfile.ZipFile(ruim, "w") as z:
        for nome, dados in partes.items():
            z.writestr(nome, dados)
    assert "o hiperlink para pg3 não tem marcador" in ed.conferir(ruim)
    assert ed.conferir(tmp_path / "nao-existe.docx")[0].startswith("o DOCX não abre")


# ----------------------------------------------------------------------
# A revisão de código (rodada 1)
# ----------------------------------------------------------------------

def test_a_fonte_dos_simbolos_vai_embutida_e_ofuscada(tmp_path):
    from core.exportar import GUID_DA_FONTE, fonte_dos_simbolos, ofuscar

    recurso = fonte_dos_simbolos("1.♘f3")
    if not recurso:
        pytest.skip("nenhuma fonte de símbolos nesta máquina")
    destino = tmp_path / "a.docx"
    ed.escrever(_livro(), destino, _meta())
    with zipfile.ZipFile(destino) as z:
        tipos = z.read("[Content_Types].xml").decode("utf-8")
        tabela = z.read("word/fontTable.xml").decode("utf-8")
        ajustes = z.read("word/settings.xml").decode("utf-8")
        fonte = z.read("word/fonts/fonte1.odttf")
        documento = z.read("word/document.xml").decode("utf-8")
    with open(recurso[1], "rb") as f:
        assert fonte == ofuscar(f.read())
    assert 'Extension="odttf"' in tipos and f'w:fontKey="{GUID_DA_FONTE}"' in tabela
    assert "<w:embedTrueTypeFonts/>" in ajustes
    assert f'w:cs="{recurso[0]}"' in documento         # a fonte nos quatro atributos
    assert ed.conferir(destino) == []


def test_sem_as_sobras_do_modelo(tmp_path):
    destino = tmp_path / "a.docx"
    ed.escrever(_livro(), destino, _meta())
    with zipfile.ZipFile(destino) as z:
        assert "docProps/thumbnail.jpeg" not in z.namelist()
        assert b"thumbnail" not in z.read("_rels/.rels")
        assert b"lastPrinted" not in z.read("docProps/core.xml")


def test_pagina_nova_desenho_e_legenda_do_diagrama_nao_revisado(tmp_path):
    livro = _livro()
    livro[1].blocos[1].revisao_pendente = True
    destino = tmp_path / "a.docx"
    ed.escrever(livro, destino, _meta())
    doc = Document(str(destino))
    paragrafos = [p for p in doc.paragraphs if p.text or p._p.xpath(".//wp:docPr")]
    quebras = [p.text or "[figura]" for p in paragrafos if p.paragraph_format.page_break_before]
    assert quebras == ["Uma seção", "A solução."]       # cada página impressa numa nova
    legenda = next(p for p in doc.paragraphs if p.style.name == "Caption")
    assert legenda.text == "pretas a jogar (informado na revisão); não revisado"
    assert doc.styles["Normal"].font.name == "Georgia"
    assert round(doc.sections[0].left_margin.cm, 1) == 2.2


def test_a_largura_do_diagrama_vem_do_tamanho_da_casa(tmp_path):
    figura = Figura(_png(), 16, 16, fen=FEN, origem="render", casas_de_largura=8.0)
    destino = tmp_path / "a.docx"
    ed.escrever([PaginaExtraida(numero=0, blocos=[figura])], destino, _meta())
    largura = Document(str(destino)).inline_shapes[0].width
    assert abs(largura.pt - 8 * ed.CORPO_DA_CASA_PT) < 0.5


def test_no_epub_a_ressalva_fica_a_vista(tmp_path):
    livro = _livro()
    livro[1].blocos[1].revisao_pendente = True
    destino = tmp_path / "a.epub"
    epub.escrever(livro, destino, _meta())
    with zipfile.ZipFile(destino) as z:
        texto = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xhtml"))
    assert "<figcaption>pretas a jogar (informado na revisão); não revisado</figcaption>" in texto
    assert epub.conferir(destino) == []